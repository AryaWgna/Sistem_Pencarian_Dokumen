import os
import sys

# Tambahkan root proyek ke sys.path
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app, db
from app.models import Dokumen
from docx import Document
from werkzeug.utils import secure_filename

app = create_app()

def generate_docs():
    with app.app_context():
        # Cari semua dokumen yang belum punya file
        docs = Dokumen.query.filter((Dokumen.file_path == None) | (Dokumen.file_path == '')).all()
        if not docs:
            print("Semua dokumen sudah memiliki file fisik.")
            return

        print(f"Menemukan {len(docs)} dokumen tanpa file fisik.")
        print("Memulai pembuatan file Microsoft Word (.docx)...")

        upload_dir = app.config['UPLOAD_FOLDER']
        if not os.path.exists(upload_dir):
            os.makedirs(upload_dir)

        count = 0
        for doc in docs:
            try:
                document = Document()
                document.add_heading(doc.judul or 'Tanpa Judul', 0)
                
                p = document.add_paragraph()
                p.add_run('Kategori: ').bold = True
                p.add_run(str(doc.kategori) + '\n')
                
                p.add_run('Jenis Surat: ').bold = True
                p.add_run(str(doc.jenis_surat or '-') + '\n')
                
                p.add_run('Nomor Dokumen: ').bold = True
                p.add_run(str(doc.nomor_dokumen or '-') + '\n')
                
                tanggal = doc.tanggal_dokumen.strftime('%Y-%m-%d') if doc.tanggal_dokumen else '-'
                p.add_run('Tanggal Dokumen: ').bold = True
                p.add_run(tanggal + '\n')
                
                document.add_heading('Isi Dokumen', level=1)
                document.add_paragraph(str(doc.isi_dokumen or '-'))
                
                safe_judul = secure_filename(doc.judul or '')
                if not safe_judul:
                    safe_judul = f"surat"
                    
                filename = f"DOC_{doc.id}_{safe_judul[:30]}.docx"
                filepath = os.path.join(upload_dir, filename)
                
                document.save(filepath)
                
                # Sesuai dengan format aplikasi: "uploads\namafile.docx"
                doc.file_path = os.path.join('uploads', filename)
                count += 1
                
                if count % 50 == 0:
                    print(f"   [Progress] {count} / {len(docs)} file berhasil dibuat...")
                    
            except Exception as e:
                print(f"Error pada dokumen ID {doc.id}: {e}")
                
        db.session.commit()
        print(f"\nSelesai! Berhasil menciptakan {count} file .docx dan mengaitkannya dengan database.")

if __name__ == '__main__':
    generate_docs()
