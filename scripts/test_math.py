from docx import Document
from lxml import etree

doc = Document()
p = doc.add_paragraph()
p.add_run('Test Equation 1 (Prior):')

math_xml1 = '''
<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <m:oMath>
    <m:r><m:t xml:space="preserve">Prior P(C) = </m:t></m:r>
    <m:f>
      <m:num><m:r><m:t>Jumlah Dokumen Kelas C</m:t></m:r></m:num>
      <m:den><m:r><m:t>Total Seluruh Dokumen (315)</m:t></m:r></m:den>
    </m:f>
  </m:oMath>
</m:oMathPara>
'''

math_xml2 = '''
<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <m:oMath>
    <m:r><m:t xml:space="preserve">P(X|C) = </m:t></m:r>
    <m:f>
      <m:num>
        <m:sSub>
          <m:e><m:r><m:t>N</m:t></m:r></m:e>
          <m:sub><m:r><m:t>xc</m:t></m:r></m:sub>
        </m:sSub>
        <m:r><m:t xml:space="preserve"> + α</m:t></m:r>
      </m:num>
      <m:den>
        <m:sSub>
          <m:e><m:r><m:t>N</m:t></m:r></m:e>
          <m:sub><m:r><m:t>c</m:t></m:r></m:sub>
        </m:sSub>
        <m:r><m:t xml:space="preserve"> + α ∙ V</m:t></m:r>
      </m:den>
    </m:f>
  </m:oMath>
</m:oMathPara>
'''

math_xml3 = '''
<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <m:oMath>
    <m:r><m:t xml:space="preserve">Posterior = \log P(C) + ∑ [ Bobot TF-IDF(X) ∙ \log P(X|C) ]</m:t></m:r>
  </m:oMath>
</m:oMathPara>
'''

math_xml4 = '''
<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <m:oMath>
    <m:r><m:t xml:space="preserve">L2 Euclidean Norm = </m:t></m:r>
    <m:rad>
      <m:pr><m:degHide m:val="1"/></m:pr>
      <m:deg/>
      <m:e><m:r><m:t xml:space="preserve">∑ (Raw_TF_IDF)²</m:t></m:r></m:e>
    </m:rad>
  </m:oMath>
</m:oMathPara>
'''

try:
    doc._element.body.append(etree.fromstring(math_xml1))
    doc._element.body.append(etree.fromstring(math_xml2))
    doc._element.body.append(etree.fromstring(math_xml3))
    doc._element.body.append(etree.fromstring(math_xml4))
    doc.save('test_math.docx')
    print('Math inserted successfully.')
except Exception as e:
    print('Error:', e)
