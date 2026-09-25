import os
import sys
import uuid
import docx

# Add project root to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import create_app, db
from app.models import Dokumen

app = create_app()

def generate_files():
    upload_folder = app.config['UPLOAD_FOLDER']
    os.makedirs(upload_folder, exist_ok=True)
    
    with app.app_context():
        docs = Dokumen.query.all()
        created_count = 0
        
        for d in docs:
            # Check if file exists
            if d.file_path and os.path.exists(d.file_path):
                continue
                
            # If not, create a new docx file
            file_name = f"dokumen_{d.id}_{uuid.uuid4().hex[:8]}.docx"
            full_path = os.path.join(upload_folder, file_name)
            
            try:
                # Create docx
                doc = docx.Document()
                doc.add_heading(d.judul, 0)
                if d.nomor_dokumen:
                    doc.add_paragraph(f"Nomor: {d.nomor_dokumen}")
                if d.kategori:
                    doc.add_paragraph(f"Kategori: {d.kategori}")
                
                doc.add_heading("Isi Dokumen:", level=1)
                doc.add_paragraph(d.isi_dokumen)
                
                doc.save(full_path)
                
                # Update DB
                d.file_path = full_path
                created_count += 1
                
                print(f"Created file for Doc ID {d.id}: {file_name}")
            except Exception as e:
                print(f"Error creating file for Doc ID {d.id}: {e}")
                
        db.session.commit()
        print(f"\nDone! Generated {created_count} physical files for documents.")

if __name__ == '__main__':
    generate_files()
